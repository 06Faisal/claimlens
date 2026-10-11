"""Turn an uploaded policy (PDF, DOCX, TXT/MD) into clauses.

Files are untrusted. Defence in depth, none of which relies on an antivirus engine:
  * allow-list by magic bytes (the filename and Content-Type are ignored)
  * the file is only ever parsed for text, never rendered, opened by an app or executed, and never stored
  * size, page, zip-entry and zip-expansion caps
  * files carrying active content (PDF scripts/launch/embedded files, Word macros/OLE/external templates,
    nested archives, executables, HTML/script text) are rejected, not cleaned
"""
import io
import re
import zipfile
from datetime import date

from .ingest import Clause

MAX_FILE = 5 * 1024 * 1024
MAX_TEXT = 300_000
MAX_CHUNKS = 400
MAX_CHUNK_LEN = 2000
MAX_PAGES = 150
MAX_ZIP_ENTRIES = 500
MAX_ZIP_TOTAL = 40 * 1024 * 1024
MAX_ZIP_RATIO = 100
CHUNK = 900

ACTIVE = ("This file contains active content (scripts, macros, links to other files or embedded files), so we cannot accept it. "
          "Please upload a plain PDF, Word (.docx) or text copy of the policy.")
UNSUPPORTED = "This file type is not supported. Please upload a PDF, Word (.docx) or text file."

# Executables, archives and OLE containers disguised as documents.
_BAD_MAGIC = (b"MZ", b"\x7fELF", b"\xfe\xed\xfa", b"\xce\xfa\xed\xfe", b"\xcf\xfa\xed\xfe", b"\xca\xfe\xba\xbe", b"#!",
              b"\x1f\x8b", b"Rar!", b"7z\xbc\xaf\x27\x1c", b"MSCF", b"{\\rtf", b"\x89PNG", b"GIF8", b"\xff\xd8\xff")
# Long, distinctive names only: short ones (/JS, /AA) also occur by chance inside compressed streams, so
# those are checked on the parsed structure instead.
_PDF_ACTIVE = (b"/JavaScript", b"/Launch", b"/EmbeddedFile", b"/OpenAction", b"/RichMedia", b"/XFA", b"/SubmitForm", b"/ImportData")
_TEXT_ACTIVE = re.compile(r"<\s*(script|iframe|object|embed|html|\?php)|javascript:|powershell|cmd\.exe|\beval\s*\(|base64,", re.I)
_DOCX_BAD_PART = re.compile(r"vbaproject|\.bin$|activex|embeddings/|oleobject|\.(exe|dll|js|vbs|bat|cmd|ps1|jar|zip|rar|7z|lnk|scr)$", re.I)


class UploadError(Exception):
    """Message is safe to show to the user."""


def _unescape_pdf_names(data: bytes) -> bytes:
    # /J#61vaScript is the same name as /JavaScript: normalise before scanning
    return re.sub(rb"#([0-9a-fA-F]{2})", lambda m: bytes([int(m.group(1), 16)]), data)


def _pdf_has_active_structure(r) -> bool:
    """Bounded check of the parsed document (catches keys hidden inside compressed object streams)."""
    root = r.trailer["/Root"]
    if "/OpenAction" in root or "/AA" in root:
        return True
    names = root.get("/Names")
    if names is not None and any(k in names.get_object() for k in ("/JavaScript", "/EmbeddedFiles")):
        return True
    form = root.get("/AcroForm")
    if form is not None and "/XFA" in form.get_object():
        return True
    for page in r.pages:
        if "/AA" in page:
            return True
        for ref in list(page.get("/Annots", []))[:2000]:
            ann = ref.get_object()
            act = ann.get("/A")
            if "/AA" in ann or ann.get("/Subtype") == "/FileAttachment":
                return True
            if act is not None and act.get_object().get("/S") in ("/JavaScript", "/Launch", "/SubmitForm", "/ImportData"):
                return True
    return False


def _pdf(data: bytes) -> str:
    from pypdf import PdfReader

    flat = _unescape_pdf_names(data)
    if any(m in flat for m in _PDF_ACTIVE):
        raise UploadError(ACTIVE)
    r = PdfReader(io.BytesIO(data))
    if r.is_encrypted:
        raise UploadError("This PDF is password-protected. Please remove the password and upload again.")
    if len(r.pages) > MAX_PAGES:
        raise UploadError(f"This PDF has more than {MAX_PAGES} pages. Please upload just the policy wording.")
    if _pdf_has_active_structure(r):
        raise UploadError(ACTIVE)
    return "\n".join((p.extract_text() or "") for p in r.pages)


def _docx(data: bytes) -> str:
    from docx import Document

    try:
        z = zipfile.ZipFile(io.BytesIO(data))
    except zipfile.BadZipFile:
        raise UploadError(UNSUPPORTED)
    with z:
        infos = z.infolist()
        if len(infos) > MAX_ZIP_ENTRIES or sum(i.file_size for i in infos) > MAX_ZIP_TOTAL:
            raise UploadError("This document is too large or complex to read.")
        names = {i.filename for i in infos}
        if "[Content_Types].xml" not in names or "word/document.xml" not in names:
            raise UploadError(UNSUPPORTED)  # some other zip (jar, apk, xlsx...) renamed to look like a document
        for i in infos:
            if i.filename.startswith("/") or ".." in i.filename.split("/"):
                raise UploadError(UNSUPPORTED)
            if i.compress_size and i.file_size / i.compress_size > MAX_ZIP_RATIO and i.file_size > 1_000_000:
                raise UploadError("This document is too large or complex to read.")
            if _DOCX_BAD_PART.search(i.filename):
                raise UploadError(ACTIVE)
        if b"macroEnabled" in z.read("[Content_Types].xml"):
            raise UploadError(ACTIVE)
        for n in names:
            if n.endswith(".rels"):  # remote templates / linked files are a known attack route
                for m in re.finditer(rb"<Relationship\b[^>]*>", z.read(n)):
                    tag = m.group(0)
                    if b'TargetMode="External"' in tag and b"/hyperlink" not in tag:
                        raise UploadError(ACTIVE)
    d = Document(io.BytesIO(data))
    parts = [p.text for p in d.paragraphs]
    for t in d.tables:
        parts += [" | ".join(c.text for c in row.cells) for row in t.rows]
    return "\n".join(parts)


def _plain(data: bytes) -> str:
    for enc in ("utf-8-sig", "cp1252"):  # cp1252 covers files saved by older Windows editors
        try:
            text = data.decode(enc)
            break
        except UnicodeDecodeError:
            continue
    else:
        raise UploadError(UNSUPPORTED)
    sample = text[:20000]
    if "\x00" in sample or sum(c.isprintable() or c in "\n\r\t" for c in sample) < 0.97 * len(sample):
        raise UploadError(UNSUPPORTED)
    if _TEXT_ACTIVE.search(text):
        raise UploadError(ACTIVE)
    return text


def extract_text(data: bytes) -> str:
    try:
        if data[:5] == b"%PDF-":
            text = _pdf(data)
        elif data[:4] == b"PK\x03\x04":
            text = _docx(data)
        elif data[:8] == b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1":
            raise UploadError("Old .doc files are not supported. Please save it as .docx or PDF and upload again.")
        elif data[:2] == b"PK" or data.startswith(_BAD_MAGIC):
            raise UploadError(UNSUPPORTED)
        else:
            text = _plain(data)
    except UploadError:
        raise
    except Exception:
        raise UploadError("We could not read this file. Please check it opens normally and try again.")
    text = re.sub(r"[^\S\n]+", " ", re.sub(r"[\x00-\x08\x0b-\x1f\x7f]", " ", text)).strip()  # drop control characters
    if len(text) < 200:
        raise UploadError("We found almost no text. If this is a scanned PDF or photo, please upload a text-based copy "
                          "(most insurers email one) or paste the text into a .txt file.")
    return text[:MAX_TEXT]


def to_chunks(text: str) -> list[str]:
    # ponytail: fixed-size paragraph chunks, no heading detection; add layout-aware sectioning if retrieval misses clauses
    out, buf = [], ""

    def flush():
        nonlocal buf
        if buf.strip():
            out.append(" ".join(buf.split()))
        buf = ""

    for para in re.split(r"\n\s*\n|\n(?=\s*\d+(?:\.\d+)*[.)]?\s+[A-Z])", text):
        while len(para) > 2 * CHUNK:  # very long block: cut at a sentence end if possible
            cut = para.rfind(". ", 0, CHUNK) + 1 or CHUNK
            if len(buf) + cut > CHUNK:
                flush()
            buf += para[:cut] + " "
            para = para[cut:]
        if len(buf) + len(para) > CHUNK:
            flush()
        buf += para + "\n"
    flush()
    return out[:MAX_CHUNKS]


def clauses_from_chunks(chunks: list[str], doc_id: str, name: str = "Your policy") -> dict[str, Clause]:
    out = {}
    for n, t in enumerate(chunks, 1):
        cid = f"{doc_id}:s{n}"
        out[cid] = Clause(cid, doc_id, "policy", date(1900, 1, 1), f"{name}, part {n}", t)
    return out
